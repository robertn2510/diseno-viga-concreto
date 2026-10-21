"""Renderizador DOCX de la memoria de cálculo (python-docx)."""

from __future__ import annotations

from pathlib import Path

from .documento import Documento

AZUL = "1F3864"
GRIS = "F2F2F2"
GRIS_TEXTO = "595959"


def _requerir_docx():
    try:
        import docx  # type: ignore
        from docx.shared import Cm, Pt, RGBColor  # type: ignore
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
        from docx.enum.table import WD_TABLE_ALIGNMENT  # type: ignore
        from docx.oxml import OxmlElement  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        return docx, Cm, Pt, RGBColor, WD_ALIGN_PARAGRAPH, WD_TABLE_ALIGNMENT, OxmlElement, qn
    except Exception as exc:  # pragma: no cover
        raise RuntimeError(
            "Para generar .docx instale python-docx: pip install python-docx. "
            "También puede generar únicamente Markdown con --formato md."
        ) from exc


def _sombreado(celda, color: str, OxmlElement, qn) -> None:
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _parrafo_sombreado(par, color: str, OxmlElement, qn) -> None:
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def _numero_pagina(par, OxmlElement, qn) -> None:
    run = par.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def renderizar_docx(doc: Documento, ruta: str | Path, encabezado: str = "") -> Path:
    docx, Cm, Pt, RGBColor, WD_ALIGN_PARAGRAPH, WD_TABLE_ALIGNMENT, OxmlElement, qn = _requerir_docx()
    ruta = Path(ruta)

    documento = docx.Document()
    documento.core_properties.title = "Memoria de cálculo — vigas de concreto armado E.060"
    documento.core_properties.subject = "Diseño estructural según RNE E.060"
    documento.core_properties.author = "vigas-e060"

    for section in documento.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)

    estilos = documento.styles
    normal = estilos["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for nombre, tam, color in (
        ("Heading 1", 16, AZUL),
        ("Heading 2", 13, AZUL),
        ("Heading 3", 11.5, "2F5597"),
        ("Heading 4", 10.5, "404040"),
    ):
        st = estilos[nombre]
        st.font.name = "Calibri"
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    if encabezado:
        header = documento.sections[0].header
        hp = header.paragraphs[0]
        hp.text = encabezado
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in hp.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(GRIS_TEXTO)

    footer = documento.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Memoria de cálculo E.060 — página ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRIS_TEXTO)
    _numero_pagina(fp, OxmlElement, qn)

    for b in doc.bloques:
        if b.tipo == "h1":
            documento.add_heading(b.texto, level=1)
        elif b.tipo == "h2":
            documento.add_heading(b.texto, level=2)
        elif b.tipo == "h3":
            documento.add_heading(b.texto, level=3)
        elif b.tipo == "h4":
            p = documento.add_paragraph()
            r = p.add_run(b.texto)
            r.bold = True
            r.italic = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor.from_string("404040")
        elif b.tipo == "p":
            p = documento.add_paragraph(b.texto)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif b.tipo == "bullets":
            for item in b.items:
                documento.add_paragraph(item, style="List Bullet")
        elif b.tipo == "tabla":
            if not b.encabezados:
                continue
            ncols = len(b.encabezados)
            tabla = documento.add_table(rows=1, cols=ncols)
            tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
            try:
                tabla.style = "Light Grid Accent 1"
            except Exception:
                tabla.style = "Table Grid"
            hdr = tabla.rows[0].cells
            for i, titulo in enumerate(b.encabezados):
                hdr[i].text = ""
                p = hdr[i].paragraphs[0]
                r = p.add_run(str(titulo))
                r.bold = True
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor.from_string("FFFFFF")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _sombreado(hdr[i], AZUL, OxmlElement, qn)
            for fila in b.filas:
                fila = list(fila) + [""] * (ncols - len(fila))
                celdas = tabla.add_row().cells
                for i, valor in enumerate(fila):
                    celdas[i].text = ""
                    p = celdas[i].paragraphs[0]
                    r = p.add_run("" if valor is None else str(valor))
                    r.font.size = Pt(8)
                    if i > 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if b.anchos:
                for fila in tabla.rows:
                    for i, ancho in enumerate(b.anchos[:ncols]):
                        fila.cells[i].width = Cm(ancho)
            documento.add_paragraph()
        elif b.tipo == "formula":
            if b.referencia:
                p = documento.add_paragraph()
                r = p.add_run(b.referencia)
                r.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor.from_string(GRIS_TEXTO)
            for linea in b.texto.splitlines():
                p = documento.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run(linea)
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                _parrafo_sombreado(p, GRIS, OxmlElement, qn)
            documento.add_paragraph()
        elif b.tipo == "nota":
            p = documento.add_paragraph()
            r = p.add_run("Nota: ")
            r.bold = True
            r.font.size = Pt(9)
            r = p.add_run(b.texto)
            r.italic = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string(GRIS_TEXTO)
            _parrafo_sombreado(p, "FFF2CC", OxmlElement, qn)
        elif b.tipo == "salto":
            documento.add_page_break()

    ruta.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(ruta))
    return ruta
