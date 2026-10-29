"""
Importación de tablas exportadas de ETABS (v20 – v23).

Formatos soportados
-------------------
* CSV / TXT exportados con  File > Export > Tables (delimitador , ; o tab).
* Excel (.xlsx) cuando `openpyxl` está instalado.
* Un solo archivo con varias tablas (encabezados "TABLE: ...") o varios
  archivos separados (fuerzas, secciones, materiales).

Tablas reconocidas
------------------
* "Element Forces - Frames"        → solicitaciones P, V2, V3, T, M2, M3
* "Frame Assignments - Section Properties"  → sección asignada por barra
* "Frame Section Properties - Rectangular" / "... - Concrete Rectangular"
                                   → dimensiones b, h por sección
* "Material Properties - Concrete Data" → f'c
"""

from __future__ import annotations

import csv
import io
import math
import os
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from ..unidades import (SistemaUnidades, fuerza_a_kgf, longitud_a_cm,
                        sistema_unidades)
from .modelo import Elemento, FuenteETABS, PuntoDemanda


# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------
def _norm(texto: str) -> str:
    if texto is None:
        return ""
    t = str(texto).strip().lower()
    t = "".join(c for c in unicodedata.normalize("NFD", t)
                if unicodedata.category(c) != "Mn")
    return t


def _sin_unidad(texto: str) -> Tuple[str, str]:
    """'P (tonf)' → ('p', 'tonf')."""
    t = _norm(texto)
    m = re.search(r"\(([^)]*)\)", t)
    unidad = m.group(1).strip() if m else ""
    nombre = re.sub(r"\([^)]*\)", "", t).strip()
    nombre = nombre.replace("'", "").replace(" ", "")
    return nombre, unidad


ALIAS = {
    "frame": {"frame", "framename", "objectid", "unique name", "uniquename",
              "element", "elem", "label", "barra", "nombre"},
    "story": {"story", "level", "piso", "storyname"},
    "station": {"station", "dist", "location", "distance", "estacion", "pos"},
    "case": {"outputcase", "loadcase", "case", "combo", "combination",
             "loadcombination", "caso", "combinacion"},
    "p": {"p", "axial", "pu"},
    "v2": {"v2", "shear2"},
    "v3": {"v3", "shear3"},
    "t": {"t", "torsion", "torsor"},
    "m2": {"m2", "moment2"},
    "m3": {"m3", "moment3"},
    "section": {"designsect", "analysectsect", "analysect", "section",
                "seccion", "propname", "frame_type"},
    "t3": {"t3", "depth", "peralte", "h", "d"},
    "t2": {"t2", "width", "ancho", "b"},
    "material": {"material", "matprop", "mat"},
    "fc": {"fc", "f'c", "fcu", "resistencia"},
}


def _campo(nombre: str) -> Optional[str]:
    n = _norm(nombre).replace(" ", "")
    for campo, alias in ALIAS.items():
        if n in alias:
            return campo
    return None


# ---------------------------------------------------------------------------
# Lectura de archivos
# ---------------------------------------------------------------------------
@dataclass
class Tabla:
    nombre: str
    filas: List[List[str]] = field(default_factory=list)

    def encabezado(self) -> List[str]:
        return self.filas[0] if self.filas else []

    def registros(self) -> List[Dict[str, str]]:
        if not self.filas:
            return []
        cab = self.filas[0]
        salida = []
        for fila in self.filas[1:]:
            if not any(str(c).strip() for c in fila):
                continue
            salida.append({cab[i]: (fila[i] if i < len(fila) else "")
                           for i in range(len(cab))})
        return salida


def _dividir_tablas(encabezado: str, filas: List[List[str]]) -> List[Tabla]:
    """Separa un archivo con varias tablas marcadas con 'TABLE: ...'."""
    tablas: List[Tabla] = []
    actual = Tabla(encabezado or "Tabla 1")
    for fila in filas:
        primera = str(fila[0]).strip() if fila else ""
        if not any(str(c).strip() for c in fila):
            continue
        if _norm(primera).startswith("table:"):
            if actual.filas:
                tablas.append(actual)
            actual = Tabla(primera[6:].strip())
            continue
        if _norm(primera).startswith("etabs") or _norm(primera).startswith("units:"):
            continue
        actual.filas.append([str(c) for c in fila])
    if actual.filas:
        tablas.append(actual)
    return tablas


def leer_tablas(ruta: str, delimitador: Optional[str] = None,
                hoja: Optional[str] = None) -> List[Tabla]:
    """Lee un CSV/TXT/XLSX y devuelve las tablas encontradas."""
    ext = os.path.splitext(ruta)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _leer_excel(ruta, hoja)
    return _leer_texto(ruta, delimitador)


def _leer_texto(ruta: str, delimitador: Optional[str]) -> List[Tabla]:
    with open(ruta, "r", encoding="utf-8-sig", errors="replace") as fh:
        muestra = fh.read(8192)
        fh.seek(0)
        if delimitador is None:
            delim = ","
            if muestra.count(";") > muestra.count(","):
                delim = ";"
            if muestra.count("\t") > max(muestra.count(","), muestra.count(";")):
                delim = "\t"
        lector = csv.reader(fh, delimiter=delim)
        filas = [fila for fila in lector]
    encabezado = ""
    for fila in filas[:5]:
        if fila and _norm(str(fila[0])).startswith("table:"):
            encabezado = str(fila[0])[6:].strip()
            break
    return _dividir_tablas(encabezado, filas)


def _leer_excel(ruta: str, hoja: Optional[str]) -> List[Tabla]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "Para leer archivos Excel instale openpyxl (pip install openpyxl) "
            "o exporte la tabla como CSV desde ETABS.") from exc
    wb = openpyxl.load_workbook(ruta, data_only=True)
    hojas = [hoja] if hoja else wb.sheetnames
    tablas = []
    for nombre in hojas:
        ws = wb[nombre]
        filas = [[("" if c is None else str(c)) for c in fila]
                 for fila in ws.iter_rows(values_only=True)]
        filas = [f for f in filas if any(str(c).strip() for c in f)]
        tablas.extend(_dividir_tablas(nombre, filas))
    return tablas


# ---------------------------------------------------------------------------
# Interpretación de tablas
# ---------------------------------------------------------------------------
def _mapear_columnas(cabecera: Sequence[str]) -> Dict[str, Tuple[int, str]]:
    mapa: Dict[str, Tuple[int, str]] = {}
    for i, col in enumerate(cabecera):
        nombre, unidad = _sin_unidad(col)
        campo = _campo(nombre)
        if campo and campo not in mapa:
            mapa[campo] = (i, unidad)
    return mapa


def _sistema_de(cabecera: Sequence[str], unidades: Optional[str]) -> SistemaUnidades:
    """Deduce el sistema de unidades a partir de los encabezados."""
    if unidades:
        return sistema_unidades(unidades)
    for col in cabecera:
        n, u = _sin_unidad(col)
        if u and ("-" in u or "_" in u):
            return sistema_unidades(u.replace("-", "_"))
    fuerzas, longitudes = set(), set()
    for col in cabecera:
        n, u = _sin_unidad(col)
        if not u:
            continue
        if u.lower() in ("tonf", "kgf", "kn", "n", "kip", "lb", "lbf", "tnf"):
            fuerzas.add(u.lower())
        if u.lower() in ("m", "cm", "mm", "in", "ft"):
            longitudes.add(u.lower())
    if len(fuerzas) == 1 and len(longitudes) == 1:
        return sistema_unidades(f"{fuerzas.pop()}_{longitudes.pop()}_C")
    return sistema_unidades("tonf_m_C")


def _a_float(valor, defecto: float = 0.0) -> float:
    try:
        t = str(valor).strip().replace(",", "")
        if t in ("", "-", "None", "nan"):
            return defecto
        return float(t)
    except (TypeError, ValueError):
        return defecto


def _tipo_por_nombre(nombre: str) -> str:
    n = _norm(nombre)
    if re.search(r"(^|[^a-z])(col|c-|c\d|columna)", n):
        return "Columna"
    if re.search(r"(^|[^a-z])(viga|vp|vs|v-|v\d|beam|b-)", n):
        return "Viga"
    return "Desconocido"


# ---------------------------------------------------------------------------
# Fuente basada en tablas
# ---------------------------------------------------------------------------
class ETABSTablas(FuenteETABS):
    """Modelo importado desde tablas exportadas de ETABS."""

    nombre = "Tablas exportadas de ETABS"
    descripcion = "Importa 'Element Forces - Frames' (y opcionalmente secciones)"

    def __init__(self, elementos: Dict[str, Elemento],
                 sistema: Optional[SistemaUnidades] = None,
                 origen: str = ""):
        self._elementos = elementos
        self.sistema = sistema or sistema_unidades("tonf_m_C")
        self.origen = origen

    # ------------------------------------------------------------------
    @classmethod
    def desde_archivos(cls, ruta_fuerzas: str, ruta_secciones: Optional[str] = None,
                       ruta_materiales: Optional[str] = None,
                       unidades: Optional[str] = None,
                       tipo_defecto: str = "auto") -> "ETABSTablas":
        tablas_f = leer_tablas(ruta_fuerzas)
        tabla_f = _buscar_tabla(tablas_f, ["element forces", "frame forces", "fuerzas"])
        if tabla_f is None:
            raise ValueError(
                "No se encontró la tabla 'Element Forces - Frames' en el archivo.\n"
                "En ETABS: Display > Tables > Analysis > Element Forces - Frames, "
                "luego File > Export > Tables > to CSV.")
        dims: Dict[str, Tuple[float, float]] = {}
        mats: Dict[str, float] = {}
        asignaciones: Dict[str, str] = {}
        if ruta_secciones:
            for t in leer_tablas(ruta_secciones):
                _extraer_dimensiones(t, dims)
                _extraer_asignaciones(t, asignaciones)
        if ruta_materiales:
            for t in leer_tablas(ruta_materiales):
                _extraer_materiales(t, mats)

        cabecera = tabla_f.encabezado()
        mapa = _mapear_columnas(cabecera)
        sistema = _sistema_de(cabecera, unidades)

        requeridos = {"frame", "station", "case", "m3"}
        faltan = requeridos - set(mapa)
        if faltan:
            raise ValueError(
                "La tabla de fuerzas no tiene las columnas esperadas: "
                + ", ".join(sorted(faltan))
                + f"\nColumnas encontradas: {cabecera}")

        grupos: Dict[str, List[PuntoDemanda]] = {}
        seccion_de: Dict[str, str] = {}
        story_de: Dict[str, str] = {}

        for reg in tabla_f.registros():
            nombre = str(reg[cabecera[mapa["frame"][0]]]).strip()
            if not nombre:
                continue
            caso = str(reg[cabecera[mapa["case"][0]]]).strip()

            def val(campo: str) -> float:
                if campo not in mapa:
                    return 0.0
                return _a_float(reg[cabecera[mapa[campo][0]]])

            est = val("station")
            p = val("p")
            v2 = val("v2")
            v3 = val("v3")
            t = val("t")
            m2 = val("m2")
            m3 = val("m3")
            unidad_est = mapa["station"][1] or sistema.longitud
            puntos = grupos.setdefault(nombre, [])
            puntos.append(PuntoDemanda(
                combo=caso,
                estacion=longitud_a_cm(est, unidad_est),
                P=sistema.fuerza_a_kgf(p),
                V2=sistema.fuerza_a_kgf(v2),
                V3=sistema.fuerza_a_kgf(v3),
                T=sistema.momento_a_kgfcm(t),
                M2=sistema.momento_a_kgfcm(m2),
                M3=sistema.momento_a_kgfcm(m3),
            ))
            if "section" in mapa:
                seccion_de[nombre] = str(reg[cabecera[mapa["section"][0]]]).strip()
            if "story" in mapa:
                story_de[nombre] = str(reg[cabecera[mapa["story"][0]]]).strip()

        # completa la sección asignada con la tabla de asignaciones (si existe)
        for nombre in grupos:
            if not seccion_de.get(nombre):
                seccion_de[nombre] = asignaciones.get(nombre, "")
        seccion_unica = list(dims)[0] if len(dims) == 1 else ""

        elementos: Dict[str, Elemento] = {}
        for nombre, pts in grupos.items():
            L = max((p.estacion for p in pts), default=0.0)
            seccion = seccion_de.get(nombre, "")
            if not seccion and seccion_unica:
                seccion = seccion_unica
            b, h = dims.get(seccion, (0.0, 0.0))
            fc = mats.get(seccion, mats.get("_default", 210.0))
            tipo = _tipo_por_nombre(nombre) if tipo_defecto == "auto" else tipo_defecto
            elementos[nombre] = Elemento(
                nombre=nombre, tipo=tipo, seccion=seccion,
                story=story_de.get(nombre, ""), material="",
                b=b, h=h, fc=fc, fy=4200.0, L=L, puntos=pts,
                unidad_fuerza=sistema.fuerza, unidad_longitud=sistema.longitud,
                origen=os.path.basename(ruta_fuerzas))
        if not elementos:
            raise ValueError("No se pudieron extraer elementos del archivo de fuerzas.")
        return cls(elementos, sistema, origen=os.path.basename(ruta_fuerzas))

    # ------------------------------------------------------------------
    def listar_elementos(self, solo_seleccionados: bool = False,
                         tipo: Optional[str] = None) -> List[str]:
        salida = sorted(self._elementos)
        if tipo and tipo != "Todos":
            salida = [n for n in salida
                      if self._elementos[n].tipo.lower().startswith(tipo.lower())]
        return salida

    def obtener_elemento(self, nombre: str) -> Elemento:
        if nombre not in self._elementos:
            raise KeyError(f"El elemento {nombre!r} no existe en el modelo importado.")
        return self._elementos[nombre]

    def listar_combos(self) -> List[str]:
        combos: List[str] = []
        for el in self._elementos.values():
            for c in el.combos:
                if c not in combos:
                    combos.append(c)
        return combos

    def aplicar_dimensiones(self, b: float, h: float, fc: float,
                            fy: float) -> None:
        """Completa dimensiones y materiales cuando no vienen en las tablas."""
        for el in self._elementos.values():
            if el.b <= 0:
                el.b = b
            if el.h <= 0:
                el.h = h
            if el.fc <= 0 or el.fc == 210.0:
                el.fc = fc
            el.fy = fy

    def info_modelo(self) -> dict:
        return {
            "fuente": self.nombre, "archivo": self.origen,
            "unidades": self.sistema.etiqueta,
            "elementos": len(self._elementos),
            "combos": len(self.listar_combos()),
        }


# ---------------------------------------------------------------------------
def _buscar_tabla(tablas: Sequence[Tabla], palabras: Sequence[str]) -> Optional[Tabla]:
    for t in tablas:
        n = _norm(t.nombre)
        if any(p in n for p in palabras) and t.filas:
            return t
    # fallback: tabla que tenga columnas de fuerzas
    for t in tablas:
        mapa = _mapear_columnas(t.encabezado())
        if {"frame", "m3", "case"} <= set(mapa):
            return t
    return None


def _extraer_dimensiones(tabla: Tabla, dims: Dict[str, Tuple[float, float]]) -> None:
    """Extrae b, h de tablas de propiedades de sección."""
    cab = tabla.encabezado()
    mapa = _mapear_columnas(cab)
    if not ({"t3", "t2"} <= set(mapa)):
        return
    sistema = _sistema_de(cab, None)
    for reg in tabla.registros():
        nombre = ""
        for clave in ("section", "frame"):
            if clave in mapa:
                nombre = str(reg[cab[mapa[clave][0]]]).strip()
                break
        if not nombre:
            continue
        u3 = mapa["t3"][1] or sistema.longitud
        u2 = mapa["t2"][1] or sistema.longitud
        h = longitud_a_cm(_a_float(reg[cab[mapa["t3"][0]]]), u3)
        b = longitud_a_cm(_a_float(reg[cab[mapa["t2"][0]]]), u2)
        if b > 0 and h > 0:
            dims[nombre] = (b, h)


def _extraer_asignaciones(tabla: Tabla, mapa_salida: Dict[str, str]) -> None:
    """
    Extrae la sección asignada a cada barra de la tabla
    "Frame Assignments - Section Properties" (Frame + DesignSect/AnalysisSect).
    """
    cab = tabla.encabezado()
    mapa = _mapear_columnas(cab)
    if "frame" not in mapa:
        return
    campo_sec = None
    for i, col in enumerate(cab):
        n = _norm(col)
        if n in ("designsect", "analysect", "analysectsect", "section", "propname"):
            campo_sec = i
            break
    if campo_sec is None:
        return
    for reg in tabla.registros():
        frame = str(reg[cab[mapa["frame"][0]]]).strip()
        seccion = str(reg[cab[campo_sec]]).strip()
        if frame and seccion and seccion.lower() not in ("none", "auto"):
            mapa_salida[frame] = seccion


def _extraer_materiales(tabla: Tabla, mats: Dict[str, float]) -> None:
    cab = tabla.encabezado()
    mapa = _mapear_columnas(cab)
    if "fc" not in mapa:
        return
    sistema = _sistema_de(cab, None)
    for reg in tabla.registros():
        nombre = ""
        if "material" in mapa:
            nombre = str(reg[cab[mapa["material"][0]]]).strip()
        fc_val = _a_float(reg[cab[mapa["fc"][0]]])
        if fc_val <= 0:
            continue
        # la tensión se convierte fuerza/longitud²
        fc = fc_val * (sistema.fuerza_a_kgf(1.0) / (sistema.longitud_a_cm(1.0) ** 2))
        if 100 < fc < 1500:
            mats[nombre or "_default"] = fc
